import csv
from datetime import timedelta
from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.contenttypes.models import ContentType
from django.core.cache import cache
from django.db import transaction
from django.test import TestCase, SimpleTestCase, override_settings
from django.utils import timezone, translation
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import Account
from assets.models import Asset, Node, Platform
from orgs.models import Organization
from orgs.utils import set_to_root_org, tmp_to_org
from perms.models import AssetPermission
from rbac.builtin import BuiltinRole
from rbac.models import Permission, Role, RoleBinding
from tickets.const import StepState
from tickets.models import Ticket, TicketAssignee, TicketStep
from users.models import User, UserGroup
from .models import UserAuthorization
from .user_authorization import (
    UserAuthorizationViewSet, authorization_rows, export_authorizations,
    permission_grantor,
)


@override_settings(AUTHENTICATION_BACKENDS=['rbac.backends.RBACBackend'])
class UserAuthorizationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        set_to_root_org()
        cls.org = Organization.objects.create(name='Authorization audit A')
        cls.other_org = Organization.objects.create(name='Authorization audit B')
        Organization.expire_orgs_mapping()
        cls.alice, cls.bob, cls.outsider, cls.auditor = User.objects.bulk_create([
            User(username=name, name=display, email=f'{name}@example.test')
            for name, display in [('audit-alice', 'Alice'), ('audit-bob', 'Bob'),
                                  ('audit-outsider', 'Outsider'), ('audit-reviewer', 'Reviewer')]
        ])
        role = Role.objects.create(name='Authorization auditor', scope='org')
        content_type = ContentType.objects.get_for_model(UserAuthorization, for_concrete_model=False)
        cls.audit_permission, _ = Permission.objects.get_or_create(
            content_type=content_type, codename='view_userauthorization',
            defaults={'name': 'Can view user authorization'},
        )
        role.permissions.add(cls.audit_permission)
        member_role = Role.objects.create(name='Authorization member', scope='org')
        RoleBinding.objects_raw.bulk_create([
            RoleBinding(user=user, org=org, role=selected_role, scope='org')
            for user, org, selected_role in [
                (cls.alice, cls.org, member_role), (cls.bob, cls.org, member_role),
                (cls.auditor, cls.org, role), (cls.outsider, cls.other_org, member_role),
            ]
        ])
        platform = Platform.objects.create(name='Audit test Linux', category='host', type='linux')
        cls.parent, child, leaf, other = Node.objects.bulk_create([
            Node(org_id=str(org.id), key=key, value=name, full_value=path)
            for org, key, name, path in [
                (cls.org, '9100', 'Production', '/Production'),
                (cls.org, '9100:1', 'Servers', '/Production/Servers'),
                (cls.org, '9100:1:1', 'Linux', '/Production/Servers/Linux'),
                (cls.other_org, '9200', 'Other', '/Other'),
            ]
        ])
        cls.asset1, cls.asset2, cls.other_asset = Asset.objects.bulk_create([
            Asset(org_id=str(org.id), name=name, address=address, platform=platform)
            for org, name, address in [(cls.org, 'web-1', '10.0.0.1'),
                                       (cls.org, 'db-1', '10.0.0.2'),
                                       (cls.other_org, 'private-1', '10.1.0.1')]
        ])
        Asset.nodes.through.objects.bulk_create([
            Asset.nodes.through(asset=asset, node=node)
            for asset, node in [(cls.asset1, leaf), (cls.asset1, child),
                                (cls.asset2, child), (cls.other_asset, other)]
        ])
        cls.group = UserGroup.objects.create(name='Audit operators', org_id=str(cls.org.id))
        User.groups.through.objects.bulk_create([
            User.groups.through(user=user, usergroup=cls.group) for user in [cls.alice, cls.bob]
        ])
        now = timezone.now()
        cls.rule, cls.second_rule, cls.expired, cls.future, cls.disabled, cls.foreign = (
            AssetPermission.objects.bulk_create([
                AssetPermission(
                    org_id=str(org.id), name=name, created_by='Rule creator',
                    date_start=start, date_expired=end, is_active=active,
                )
                for org, name, start, end, active in [
                    (cls.org, 'Node rule', now - timedelta(days=1), now + timedelta(days=1), True),
                    (cls.org, 'Direct rule', now - timedelta(days=1), now + timedelta(days=1), True),
                    (cls.org, 'Expired', now - timedelta(days=2), now - timedelta(days=1), True),
                    (cls.org, 'Future', now + timedelta(days=1), now + timedelta(days=2), True),
                    (cls.org, 'Disabled', now - timedelta(days=1), now + timedelta(days=1), False),
                    (cls.other_org, 'Foreign rule', now - timedelta(days=1), now + timedelta(days=1), True),
                ]
            ])
        )
        AssetPermission.users.through.objects.bulk_create([
            AssetPermission.users.through(assetpermission=rule, user=user)
            for rule, user in [(cls.rule, cls.alice), (cls.second_rule, cls.alice),
                               (cls.expired, cls.alice), (cls.future, cls.alice),
                               (cls.disabled, cls.alice), (cls.foreign, cls.outsider)]
        ])
        AssetPermission.user_groups.through.objects.create(assetpermission=cls.rule, usergroup=cls.group)
        AssetPermission.nodes.through.objects.create(assetpermission=cls.rule, node=cls.parent)
        AssetPermission.assets.through.objects.bulk_create([
            AssetPermission.assets.through(assetpermission=rule, asset=asset)
            for rule, asset in [(cls.rule, cls.asset1), (cls.second_rule, cls.asset1),
                                (cls.expired, cls.asset1), (cls.future, cls.asset1),
                                (cls.disabled, cls.asset1), (cls.foreign, cls.other_asset)]
        ])

    def setUp(self):
        cache.clear()

    def tearDown(self):
        set_to_root_org()
        Organization.expire_orgs_mapping()

    def request(self, action='list', params=None, user=None, pk=None, method='get'):
        request = getattr(APIRequestFactory(), method)('/api/v1/audits/user-authorizations/', params or {})
        force_authenticate(request, user=user or self.auditor)
        view = UserAuthorizationViewSet.as_view({method: action}, throttle_classes=[])
        with transaction.atomic():
            return view(request, **({'pk': str(pk)} if pk else {}))

    def test_user_search_pagination_and_group_inheritance(self):
        with tmp_to_org(self.org):
            response = self.request(params={'limit': 1, 'q': 'audit-'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.data['count'], 2)
            response = self.request(params={'name__icontains': 'Bob'})
            self.assertEqual([u['username'] for u in response.data['results']], ['audit-bob'])
            response = self.request('rules', pk=self.bob.id)
            self.assertEqual([r['name'] for r in response.data['results']], ['Node rule'])

    def test_rows_include_descendants_and_keep_distinct_rules(self):
        with tmp_to_org(self.org):
            rows = list(authorization_rows(User.get_org_users()))
        self.assertEqual(len(rows), 5)
        self.assertEqual(sum(row[0] == 'audit-alice' for row in rows), 3)
        self.assertEqual({row[2] for row in rows}, {'Node rule', 'Direct rule'})
        self.assertEqual({row[5] for row in rows}, {'web-1', 'db-1'})
        self.assertEqual(next(row[7] for row in rows if row[5] == 'web-1'),
                         '/Production/Servers; /Production/Servers/Linux')

    def test_detail_asset_pagination_and_organization_boundaries(self):
        with tmp_to_org(self.org):
            response = self.request('assets', {'rule_id': str(self.rule.id), 'limit': 1}, pk=self.alice.id)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.data['count'], 2)
            self.assertEqual(len(response.data['results']), 1)
            self.assertEqual(self.request('rules', pk=self.outsider.id).status_code, 404)
            self.assertEqual(self.request('assets', {'rule_id': str(self.foreign.id)}, pk=self.alice.id).status_code, 404)
            self.assertEqual(self.request('assets', {'rule_id': 'invalid'}, pk=self.alice.id).status_code, 400)

    def test_auditor_permission_is_required(self):
        with tmp_to_org(self.org):
            for action, params, pk in [('list', {}, None), ('rules', {}, self.alice.id),
                                       ('assets', {'rule_id': str(self.rule.id)}, self.alice.id),
                                       ('list', {'format': 'xlsx'}, None)]:
                self.assertEqual(self.request(action, params, user=self.alice, pk=pk).status_code, 403)
            self.assertIn(self.audit_permission.pk, BuiltinRole.org_auditor.default_perms)

    def test_export_selected_user_does_not_use_list_pagination(self):
        with tmp_to_org(self.org), patch('audits.user_authorization.LogMixin.record_logs') as log:
            with patch('common.drf.filters.IDSpmFilterBackend.get_resource_ids', return_value=[str(self.bob.id)]):
                response = self.request(params={'format': 'csv', 'spm': 'selection', 'limit': 1})
            self.assertEqual(response.status_code, 200)
            self.addCleanup(response.file_to_stream.close)
            rows = list(csv.reader(StringIO(b''.join(response.streaming_content).decode('utf-8-sig'))))
            self.assertEqual(len(rows), 3)
            self.assertTrue(all(row[0] == 'audit-bob' for row in rows[1:]))
            self.assertEqual(len(log.call_args.args[2]), 2)

    def test_global_export_keeps_rules_in_their_organization(self):
        set_to_root_org()
        rows = list(authorization_rows(User.objects.filter(id__in=[self.alice.id, self.outsider.id])))
        self.assertEqual(len(rows), 4)
        self.assertEqual([r[5] for r in rows if r[0] == 'audit-outsider'], ['private-1'])

    def test_export_summarizes_available_accounts_per_user_rule_and_asset(self):
        Account.objects.bulk_create([
            Account(org_id=asset.org_id, asset=asset, name=name, username=username,
                    secret_type=secret_type, is_active=active)
            for asset, name, username, secret_type, active in [
                (self.asset1, 'root', 'root', 'password', True),
                (self.asset1, 'operator password', 'operator', 'password', True),
                (self.asset1, 'operator key', 'operator', 'ssh_key', True),
                (self.asset1, 'alice', self.alice.username, 'password', True),
                (self.asset1, 'disabled', 'disabled', 'password', False),
                (self.asset2, 'database', 'dbadmin', 'password', True),
                (self.other_asset, 'private', 'private-admin', 'password', True),
            ]
        ])
        AssetPermission.objects.filter(pk=self.rule.pk).update(
            accounts=['@ALL', '!root', '@USER', '@INPUT', '@ANON']
        )
        AssetPermission.objects.filter(pk=self.second_rule.pk).update(
            accounts=['root', 'missing', 'disabled']
        )
        AssetPermission.objects.filter(pk=self.foreign.pk).update(accounts=['@ALL'])
        with tmp_to_org(self.org), translation.override('en'):
            rows = list(authorization_rows(User.get_org_users()))
            summaries = {(row[0], row[2], row[5]): row[8] for row in rows}
            special = {'@INPUT (Manual input)', '@ANON (Anonymous account)'}
            self.assertEqual(summaries[('audit-alice', 'Direct rule', 'web-1')], 'root')
            self.assertEqual(set(summaries[('audit-alice', 'Node rule', 'web-1')].split('; ')),
                             special | {'audit-alice', 'operator'})
            self.assertEqual(set(summaries[('audit-bob', 'Node rule', 'web-1')].split('; ')),
                             special | {'audit-alice', 'audit-bob', 'operator'})
            self.assertEqual(set(summaries[('audit-bob', 'Node rule', 'db-1')].split('; ')),
                             special | {'audit-bob', 'dbadmin'})
            self.assertEqual(summaries[('audit-alice', 'Node rule', 'web-1')].count('audit-alice'), 1)
            self.assertEqual(summaries[('audit-alice', 'Node rule', 'web-1')].count('operator'), 1)
            for format in ('csv', 'xlsx'):
                with self.subTest(format=format), patch('audits.user_authorization.LogMixin.record_logs'):
                    response = self.request(params={'format': format, 'limit': 1})
                    self.assertEqual(response.status_code, 200)
                    self.addCleanup(response.file_to_stream.close)
                    data = b''.join(response.streaming_content)
                    if format == 'csv':
                        exported = list(csv.reader(StringIO(data.decode('utf-8-sig'))))
                    else:
                        workbook = load_workbook(BytesIO(data), read_only=True)
                        exported = [list(row) for row in workbook.active.values]
                        workbook.close()
                    self.assertEqual(exported[0][-1], 'Accounts available to user')
                    self.assertEqual(len(exported), 6)
                    for row in exported[1:]:
                        self.assertEqual(row[8].lstrip("'"), summaries[tuple(row[i] for i in (0, 2, 5))])
        with translation.override('en'):
            rows = list(authorization_rows(User.objects.filter(id=self.outsider.id)))
        self.assertEqual(rows[0][8], 'private-admin')

    def test_unavailable_accounts_leave_summary_empty(self):
        with tmp_to_org(self.org):
            AssetPermission.objects.filter(pk=self.rule.pk).update(accounts=['missing'])
            rows = list(authorization_rows(User.get_org_users()))
        self.assertTrue(rows)
        self.assertTrue(all(row[8] == '' for row in rows))

    def test_ticket_grantor_uses_only_actual_approvers(self):
        ticket = Ticket.objects.create(org_id=str(self.org.id), title='Asset request', applicant=self.bob)
        step = TicketStep.objects.create(ticket=ticket, state=StepState.approved)
        TicketAssignee.objects.bulk_create([
            TicketAssignee(step=step, assignee=self.auditor, state=StepState.approved),
            TicketAssignee(step=step, assignee=self.alice, state=StepState.pending),
        ])
        rule = AssetPermission(id=ticket.id, org_id=str(self.org.id), from_ticket=True,
                               created_by=f'ApplyAssetTicket:{ticket.id}')
        self.assertEqual(permission_grantor(rule), str(self.auditor))

    def test_options_exposes_user_and_detail_fields(self):
        with tmp_to_org(self.org):
            for action, pk, expected in [('list', None, 'username'), ('rules', self.alice.id, 'grantor'),
                                         ('assets', self.alice.id, 'asset_groups')]:
                request = APIRequestFactory().options('/api/v1/audits/user-authorizations/')
                force_authenticate(request, user=self.auditor)
                view = UserAuthorizationViewSet.as_view({'get': action}, throttle_classes=[])
                response = view(request, **({'pk': str(pk)} if pk else {}))
                self.assertEqual(response.status_code, 200)
                self.assertIn(expected, response.data['actions']['GET'])


class AuthorizationExportTests(SimpleTestCase):
    def read_export(self, rows, format):
        response, count = export_authorizations(rows, format)
        data = b''.join(response.streaming_content)
        response.close()
        return data, count

    def test_exports_all_rows_beyond_generic_limit(self):
        row = ['alice', 'creator', 'rule', 'start', 'end', 'asset', '10.0.0.1', '/Production', 'root; operator']
        for format in ['csv', 'xlsx']:
            data, count = self.read_export((row for _ in range(10005)), format)
            self.assertEqual(count, 10005)
            if format == 'csv':
                rows = list(csv.reader(StringIO(data.decode('utf-8-sig'))))
            else:
                workbook = load_workbook(BytesIO(data), read_only=True)
                rows = list(workbook.active.values)
                workbook.close()
            self.assertEqual(len(rows), 10006)
            self.assertEqual(len(rows[0]), 9)
            self.assertEqual(rows[-1][8], 'root; operator')

    def test_formula_content_is_text_and_excel_splits_at_sheet_limit(self):
        row = ['=1+1', '+creator', '@rule', 'start', 'end', '-asset', '10.0.0.1', '/Production', '=account']
        data, _ = self.read_export([row], 'csv')
        rows = list(csv.reader(StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(rows[1][0], "'=1+1")
        self.assertEqual(rows[1][8], "'=account")
        with patch('audits.user_authorization.EXCEL_SHEET_ROWS', 3):
            data, count = self.read_export([row] * 3, 'xlsx')
        workbook = load_workbook(BytesIO(data))
        self.assertEqual(count, 3)
        self.assertEqual(len(workbook.worksheets), 2)
        self.assertEqual(workbook.active['A2'].data_type, 's')
        self.assertEqual(workbook.active['A2'].value, '=1+1')
        self.assertEqual(workbook.active['I2'].data_type, 's')
        self.assertEqual(workbook.active['I2'].value, '=account')
        workbook.close()

    def test_empty_export_still_contains_headers(self):
        with translation.override('en'):
            data, count = self.read_export([], 'csv')
        self.assertEqual(count, 0)
        self.assertEqual(next(csv.reader(StringIO(data.decode('utf-8-sig')))), [
            'Username', 'Granted by', 'Authorization rule name', 'Authorization start time',
            'Authorization expiration time', 'Authorized asset', 'Asset IP', 'Asset groups',
            'Accounts available to user',
        ])
